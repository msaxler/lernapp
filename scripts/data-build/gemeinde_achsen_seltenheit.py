"""Gemeinde-Achsen: Seltenheits-Karten (Lauf s0.1; Spielkonzept v0.2.8 §8 Nr. 7).

Idee Mike (2026-10-03): „Waldbreitbach hat einen Wert, den nur 73 der 10.753 Gemeinden Deutschlands übertreffen – welchen?“
Die Seltenheit steht in der Frage, das Erstaunen kommt mit der Eigenschaft. Mikes Urteil über die Probe: Waldbreitbach
(Frauenanteil) „OK“, Horben (AfD-Tiefstwert) und Kirchzarten (Grünen-Höchstwert) „zu schwach“ – dort erklärt die Gegend
den Wert. Deshalb trägt eine Karte nur, wenn
  - höchstens ein Prozent der Gemeinden Deutschlands den Wert übertrifft (bzw. unterbietet), und
  - der Ort auch in seinem Kreis Platz eins hat, mit Abstand zum Zweiten (Anteile: ein Prozentpunkt, sonst zwei Prozent).
Vergleichsmengen: Strukturdaten aus dem Gemeindeverzeichnis des Statistischen Bundesamts (alle bewohnten Gemeinden);
Wahl aus der Wahlbezirksstatistik, nur Gemeinden mit vollständigem Ergebnis aus Urne und Brief, CDU ohne Bayern. Für die
Wahl gilt der Kreis-Vergleich nur in Kreisen, deren Gemeinden alle ein vollständiges Ergebnis haben.
Ein Satz aus den Fakten des Orts kann die Rückseite ergänzen (<raum>/seltenheit-zusatz.json, mit Beleg).

Ausgabe: <raum>/karten-s0.1/<slug>.md (nur Orte mit Karte), <raum>/seltenheit.json.
Aufruf:  python -X utf8 scripts/data-build/gemeinde_achsen_seltenheit.py      (GA_RAUM=<raum>)
Braucht data/raw/gv_auszug.xlsx (Destatis, AuszugGV 4. Quartal 2025, nicht im Repo):
https://www.destatis.de/DE/Themen/Laender-Regionen/Regionales/Gemeindeverzeichnis/Administrativ/Archiv/GVAuszugQ/AuszugGV4QAktuell.xlsx
"""
import csv
import io
import json
import os
import re
import sys
import zipfile

import openpyxl

WURZEL = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
ITER = os.path.join(WURZEL, 'data', 'gemeinde-achsen', os.environ.get('GA_RAUM', 'iter1'))
AUS = os.path.join(ITER, 'karten-s0.1')
ANTEIL_MAX, ABSTAND_ANTEIL, ABSTAND_REL = 0.01, 1.0, 0.02
# Merkmal: (Familie der Vergleichsmenge, Anteil?, Option, Wertsatz, Einheit)
MERKMALE = {
    'frauenanteil': ('gv', True, 'Den Frauenanteil', None, 'Prozent'),
    'dichte': ('gv', False, 'Die Einwohner je Quadratkilometer', '%s kommt auf %s Einwohner je Quadratkilometer.', ''),
    'flaeche': ('gv', False, 'Die Fläche der Gemarkung', '%s hat eine Gemarkung von %s Quadratkilometern.', ''),
    'einwohner': ('gv', False, 'Die Einwohnerzahl', '%s hat %s Einwohner.', ''),
    'wahlbeteiligung': ('wahl', True, 'Die Wahlbeteiligung', '%s kam auf eine Wahlbeteiligung von %s Prozent.', ''),
    'CDU': ('wahl', True, 'Den Anteil der CDU', 'Die CDU holte in %s %s Prozent.', ''),
    'AfD': ('wahl', True, 'Den Anteil der AfD', 'Die AfD holte in %s %s Prozent.', ''),
    'SPD': ('wahl', True, 'Den Anteil der SPD', 'Die SPD holte in %s %s Prozent.', ''),
    'GRÜNE': ('wahl', True, 'Den Anteil der Grünen', 'Die Grünen holten in %s %s Prozent.', ''),
    'Die Linke': ('wahl', True, 'Den Anteil der Linken', 'Die Linke holte in %s %s Prozent.', ''),
}
QUELLE = {'gv': 'Quelle: Statistisches Bundesamt, Gemeindeverzeichnis (Bevölkerung am 31.12.2024)',
          'wahl': 'Quelle: Die Bundeswahlleiterin, Wahlbezirksstatistik Bundestagswahl 2025 (Stand 23.02.2025)'}


def de(x, nach=1):
    if isinstance(x, int) or float(x).is_integer() and nach == 0:
        return '{:,}'.format(int(x)).replace(',', '.')
    return ('%.*f' % (nach, x)).replace('.', ',')


def gv_laden():
    wb = openpyxl.load_workbook(os.path.join(WURZEL, 'data', 'raw', 'gv_auszug.xlsx'), read_only=True, data_only=True)
    aus = {}
    for r in wb[wb.sheetnames[1]].iter_rows(values_only=True):
        if r[0] == '60' and isinstance(r[9], (int, float)) and r[8] and r[9]:
            aus[r[2] + r[3] + r[4] + r[6]] = {'name': r[7], 'einwohner': int(r[9]), 'flaeche': float(r[8]),
                                              'dichte': r[9] / r[8], 'frauenanteil': 100.0 * r[11] / r[9], 'frauen': int(r[11])}
    return aus


def wahl_laden():
    z = zipfile.ZipFile(os.path.join(WURZEL, 'data', 'raw', 'btw25_wbz.zip'))
    zl = list(csv.reader(io.StringIO(z.read('btw25_wbz_ergebnisse.csv').decode('utf-8-sig')), delimiter=';'))
    k = zl[4]
    ix = {c: i for i, c in enumerate(k)}
    n = lambda x: int(x) if x.strip().isdigit() else 0
    s, fremd = {}, set()
    for r in zl[5:]:
        if len(r) < len(k) or not r[ix['Land']].isdigit():
            continue
        a = r[ix['Land']] + r[ix['Regierungsbezirk']] + r[ix['Kreis']] + r[ix['Gemeinde']]
        if r[ix['Kennziffer Briefwahlzugehörigkeit']].strip() not in ('', '00'):
            fremd.add(a)
        g = s.setdefault(a, {'ber': 0, 'wae': 0, 'gue': 0, 'p': {}})
        g['ber'] += n(r[ix['Wahlberechtigte (A)']])
        g['wae'] += n(r[ix['Wählende (B)']])
        g['gue'] += n(r[ix['Gültige - Zweitstimmen']])
        for p in ('CDU', 'AfD', 'SPD', 'GRÜNE', 'Die Linke'):
            g['p'][p] = g['p'].get(p, 0) + n(r[ix[p + ' - Zweitstimmen']])
    aus = {}
    for a, g in s.items():
        if a in fremd or not g['gue'] or not g['ber']:
            continue
        aus[a] = dict({'wahlbeteiligung': 100.0 * g['wae'] / g['ber']}, **{p: 100.0 * v / g['gue'] for p, v in g['p'].items()})
    kreise_unvollstaendig = set(a[:5] for a in fremd)
    return aus, kreise_unvollstaendig


def menge(daten, merkmal):
    """Vergleichsmenge eines Merkmals: (ags, wert); CDU ohne Bayern (dort tritt die CSU an)."""
    return [(a, d[merkmal]) for a, d in daten.items() if merkmal in d and not (merkmal == 'CDU' and a.startswith('09'))]


def zaehle(m, ags, wert, hoch):
    return sum(1 for a, w in m if (w > wert if hoch else w < wert))


def median(werte):
    v = sorted(werte)
    return v[len(v) // 2]


def main():
    gv = gv_laden()
    wahl, unvoll = wahl_laden()
    daten = {'gv': gv, 'wahl': wahl}
    zusatz = {}
    pfad_z = os.path.join(ITER, 'seltenheit-zusatz.json')
    if os.path.exists(pfad_z):
        with io.open(pfad_z, encoding='utf-8') as f:
            zusatz = json.load(f)['orte']
    with io.open(os.path.join(ITER, 'grunddaten.json'), encoding='utf-8') as f:
        orte = [o for o in json.load(f)['orte'] if o['rolle'] == 'ziel' and not o['einheit'].startswith('Ortsteil')]
    os.makedirs(AUS, exist_ok=True)
    uebersicht = []
    for i, o in enumerate(orte):
        ags = o['gemeindeschluessel'][:8]
        funde = []
        for mk, (fam, anteil, _, _, _) in MERKMALE.items():
            d = daten[fam]
            if ags not in d or mk not in d[ags] or (fam == 'wahl' and ags[:5] in unvoll):
                continue
            m = menge(d, mk)
            wert = d[ags][mk]
            for hoch in (True, False):
                n = zaehle(m, ags, wert, hoch)
                if n > ANTEIL_MAX * len(m):
                    continue
                kreis = sorted((w for a, w in m if a[:5] == ags[:5] and a != ags), reverse=hoch)
                if not kreis or (hoch and kreis[0] >= wert) or (not hoch and kreis[0] <= wert):
                    continue
                abstand = abs(wert - kreis[0])
                if abstand < (ABSTAND_ANTEIL if anteil else ABSTAND_REL * abs(wert)):
                    continue
                funde.append((n / len(m), mk, fam, hoch, n, len(m), wert, kreis[0], median([w for _, w in m])))
        eintrag = {'slug': o['slug'], 'ort': o['name'], 'karten': []}
        pfad = os.path.join(AUS, o['slug'] + '.md')
        if funde:
            funde.sort()
            _, mk, fam, hoch, n, gesamt, wert, zweiter, med = funde[0]
            richtung = 'übertreffen' if hoch else 'unterbieten'
            andere = []
            for mk2, (fam2, _, _, _, _) in MERKMALE.items():
                if mk2 == mk or fam2 != fam or mk2 not in daten[fam].get(ags, {}):
                    continue
                n2 = zaehle(menge(daten[fam], mk2), ags, daten[fam][ags][mk2], hoch)
                if n2 >= 3 * max(n, 10):
                    andere.append((-n2, mk2, n2))
            andere.sort()
            opts = [x[1] for x in andere[:3]]
            pos = ((i * 3 + 7) % 4) + 1
            opts.insert(pos - 1, mk)
            menge_txt = ('der %s Gemeinden Deutschlands' % de(gesamt, 0)) if fam == 'gv' else (
                'der %s Gemeinden mit vollständigem Ergebnis' % de(gesamt, 0))
            frage = '%s%s hat einen Wert, den nur %d %s %s. Welchen?' % (
                'Bundestagswahl 2025: ' if fam == 'wahl' else '', o['name'], n, menge_txt, richtung)
            ew = int(round(o['eigenschaften']['einwohner'][0]['wert'], -2))
            lage = o['eigenschaften']['lage'][0]['wert'].split(' (')[0]
            steck = '%s mit rund %s Einwohnern, %s.' % ('Stadt' if o['einheit'] == 'Stadt' else 'Gemeinde', de(ew, 0), lage)
            g = gv[ags]
            if mk == 'frauenanteil':
                wertsatz = 'Von %s Einwohnern sind %s Frauen, %s Prozent.' % (de(g['einwohner'], 0), de(g['frauen'], 0), de(wert))
            else:
                wertsatz = MERKMALE[mk][3] % (o['name'], de(wert, 1 if MERKMALE[mk][1] else 0))
            kreisname = o['eigenschaften']['landkreis'][0]['wert']
            rueck = ('Du hattest [Option] getippt. %s Das %s nur %d %s; der Median liegt bei %s. Im Landkreis %s kommt keine '
                     'andere Gemeinde %s %s.%s Richtig war %d.' % (
                         wertsatz, richtung, n, menge_txt, de(med, 1), kreisname, 'über' if hoch else 'unter', de(zweiter, 1),
                         (' ' + zusatz[o['slug']]['satz']) if o['slug'] in zusatz else '', pos))
            nachweise = []
            for j, q in enumerate(opts):
                if q != mk:
                    nn = zaehle(menge(daten[fam], q), ags, daten[fam][ags][q], hoch)
                    nachweise.append('%d: Weg (a). %s: %s Gemeinden %s den Wert von %s, nicht %d (vollständige Quelle).' % (
                        j + 1, MERKMALE[q][2], de(nn, 0), richtung, o['name'], n))
            gewaehlt = '%s %s %s: nur %d von %d Gemeinden %s ihn; Kreis Platz eins vor %s; Median %s' % (
                o['name'], mk, de(wert, 2), n, gesamt, richtung, de(zweiter, 2), de(med, 2))
            text = '\n'.join(['## Karte 1', 'SORTE: Klassiker', 'FAMILIE: A', 'FAKT-ID: GRUNDDATEN.seltenheit_%s' % re.sub(r'\W', '', mk.lower()),
                              'GEWÄHLTER FAKT: GRUNDDATEN seltenheit_%s: %s' % (re.sub(r'\W', '', mk.lower()), gewaehlt),
                              'BEKANNTHEIT: niedrig. Den Wert kennt kaum jemand; überraschend ist, welche Eigenschaft es ist.',
                              'VORDERSEITE:', o['name'], steck, frage] +
                             ['%d. %s' % (j + 1, MERKMALE[q][2]) for j, q in enumerate(opts)] +
                             ['RÜCKSEITE:', rueck, QUELLE[fam], 'WÖRTER RÜCKSEITE: %d' % len(rueck.split()), 'NEGATIVNACHWEISE:'] + nachweise +
                             ['PRÜFHINWEIS: Zählung nachrechnen; Zusatzsatz: %s' % (zusatz.get(o['slug'], {}).get('beleg', 'keiner')), '',
                              '## ORTS-ANSCHLUSS', 'keiner (Skript-Lauf)', '## NICHT VERWENDET', 'keiner', '## UNGEREGELT', 'keiner', ''])
            with io.open(pfad, 'w', encoding='utf-8', newline='\n') as f:
                f.write(text)
            eintrag['karten'].append({'merkmal': mk, 'wert': round(wert, 2), 'gemeinden_besser': n, 'menge': gesamt,
                                      'richtung': richtung, 'kreis_zweiter': round(zweiter, 2)})
        elif os.path.exists(pfad):
            os.remove(pfad)
        uebersicht.append(eintrag)
        print('%-18s %s' % (o['slug'], '; '.join('%s %s nur %d/%d' % (c['merkmal'], c['richtung'], c['gemeinden_besser'], c['menge'])
                                                 for c in eintrag['karten']) or '-'))
    with io.open(os.path.join(ITER, 'seltenheit.json'), 'w', encoding='utf-8', newline='\n') as f:
        json.dump({'_doc': 'Seltenheits-Karten je Zielgemeinde (gemeinde_achsen_seltenheit.py, Lauf s0.1)', 'orte': uebersicht},
                  f, ensure_ascii=False, indent=1)


if __name__ == '__main__':
    sys.stdout.reconfigure(encoding='utf-8')
    main()
